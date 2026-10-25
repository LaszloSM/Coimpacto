"""Módulos del panel (Ahorro, Crédito, Cartera, Análisis, Reportes) y gestión de comunidades."""

import datetime
import io

import pytest
from django.test import Client
from django.utils import timezone
from openpyxl import load_workbook

from acap import analitica, tablero
from acap.models import Acap, Cargo, Ciclo, Membresia
from cuentas.models import Titular, Usuario
from programas.models import Proyecto
from sync.tests.test_tablero import reunion_con_prestamo

pytestmark = pytest.mark.django_db


@pytest.fixture
def admin_cliente(db):
    cliente = Client()
    cliente.force_login(Usuario.objects.create_superuser(username="adm", password="x" * 16))
    return cliente


def enviar(cliente, ruta, datos):
    return cliente.post(ruta, datos, content_type="application/json", secure=True)


def test_ahorro_por_comunidad_y_mes(junta):
    reunion_con_prestamo(junta, timezone.localdate())
    datos = analitica.ahorro()
    assert datos["resumen"]["ahorro_ciclo"] == 30_000
    [fila] = datos["comunidades"]
    assert fila["ahorro"] == 30_000
    assert fila["reuniones"] == 1
    assert datos["mensual"][-1]["ahorro"] == 30_000


def test_credito_y_cartera_con_mora(junta):
    reunion_con_prestamo(junta, timezone.localdate() - datetime.timedelta(days=75))
    credito = analitica.credito()
    assert credito["resumen"]["prestamos"] == 1
    assert credito["resumen"]["desembolsado"] == 100_000
    cartera = analitica.cartera()
    assert cartera["resumen"]["mora"] == 100.0
    [prestamo] = cartera["prestamos"]
    assert prestamo["dias_atraso"] > 0
    assert prestamo["miembro"] is not None
    tramos = {t["tramo"]: t["prestamos"] for t in cartera["tramos"]}
    assert tramos["al_dia"] == 0
    assert sum(tramos.values()) == 1
    lecturas = analitica.analisis()["lecturas"]
    assert any(lectura["tono"] == "alerta" for lectura in lecturas)


def test_reporte_en_excel(junta, admin_cliente):
    reunion_con_prestamo(junta, timezone.localdate())
    r = admin_cliente.get("/api/oficina/reportes/consolidado.xlsx", secure=True)
    assert r.status_code == 200
    libro = load_workbook(io.BytesIO(r.content))
    assert {"Resumen", "Ahorro por ACAP", "Préstamos vigentes", "Reuniones"} <= set(
        libro.sheetnames
    )


def test_modulos_piden_rol(junta):
    cliente = Client()
    cliente.force_login(Usuario.objects.create_user(username="x", password="x" * 16))
    for ruta in ("ahorro", "credito", "cartera", "analisis", "comunidades", "proyectos"):
        assert cliente.get(f"/api/oficina/{ruta}", secure=True).status_code == 403


def test_gestion_completa_desde_el_panel(admin_cliente):
    r = enviar(admin_cliente, "/api/oficina/comunidades", {"nombre": "ACAP Nueva", "zona": "Norte"})
    assert r.status_code == 200, r.content
    acap_id = r.json()["id"]
    assert (
        enviar(admin_cliente, "/api/oficina/comunidades", {"nombre": "acap nueva"}).status_code
        == 400
    )

    r = enviar(
        admin_cliente,
        f"/api/oficina/comunidades/{acap_id}/ciclos",
        {
            "fecha_inicio": "2026-10-01",
            "duracion_meses": 12,
            "valor_accion": 10000,
            "valor_fondo_social": 2000,
            "valor_multa": 2000,
            "tasa_interes": "3",
        },
    )
    assert r.status_code == 200, r.content
    ciclo = Ciclo.objects.get(acap_id=acap_id)
    assert ciclo.estado == Ciclo.Estado.ACTIVO
    assert ciclo.numero == 1

    membresia = Membresia.objects.create(
        acap_id=acap_id, titular=Titular.objects.create(), numero_miembro=1
    )
    ruta_cargos = f"/api/oficina/ciclos/{ciclo.id}/cargos"
    r = enviar(
        admin_cliente, ruta_cargos, {"tipo": "presidente", "membresia_id": str(membresia.id)}
    )
    assert r.status_code == 200
    r = enviar(
        admin_cliente, ruta_cargos, {"tipo": "portador_caja", "membresia_id": str(membresia.id)}
    )
    vigentes = Cargo.objects.filter(ciclo=ciclo, vigente=True).values_list("tipo", flat=True)
    assert list(vigentes) == ["portador_caja"]

    ruta_usuario = f"/api/oficina/comunidades/{acap_id}/usuario-app"
    assert (
        enviar(admin_cliente, ruta_usuario, {"usuario": "nueva", "clave": "corta"}).status_code
        == 400
    )
    r = enviar(admin_cliente, ruta_usuario, {"usuario": "nueva", "clave": "una-clave-larga-1"})
    assert r.status_code == 200
    assert Usuario.objects.get(username="nueva").acap_id == Acap.objects.get(id=acap_id).id

    detalle = admin_cliente.get(f"/api/oficina/comunidades/{acap_id}", secure=True).json()
    assert detalle["usuarios_app"][0]["usuario"] == "nueva"
    assert detalle["cargos"][0]["tipo"] == "portador_caja"


def test_proyectos_capacitaciones_y_evaluaciones(admin_cliente):
    acap = Acap.objects.create(nombre="ACAP Programas")
    r = enviar(
        admin_cliente,
        "/api/oficina/proyectos",
        {"acap_id": str(acap.id), "nombre": "Huerta", "estado": "ejecucion", "avance": 40},
    )
    assert r.status_code == 200, r.content
    malo = {"acap_id": str(acap.id), "nombre": "X", "avance": 140}
    assert enviar(admin_cliente, "/api/oficina/proyectos", malo).status_code == 400
    r = enviar(
        admin_cliente,
        "/api/oficina/capacitaciones",
        {"tema": "app", "titulo": "Uso de la app", "fecha": "2026-10-20", "asistentes": 12},
    )
    assert r.status_code == 200
    r = enviar(
        admin_cliente,
        "/api/oficina/evaluaciones",
        {
            "acap_id": str(acap.id),
            "fecha": "2026-10-20",
            "gobernanza": 4,
            "transparencia": 5,
            "participacion": 3,
            "gestion_financiera": 4,
            "reglamento": 4,
        },
    )
    assert r.json()["puntaje"] == 4.0
    indicadores = tablero.tablero()["indicadores"]
    assert indicadores["proyectos"] == 1
    assert indicadores["capacitaciones"] == 1
    assert Proyecto.objects.get().avance == 40


def test_vista_previa_del_cierre_de_ciclo(junta, admin_cliente):
    reunion_con_prestamo(junta, timezone.localdate())
    r = admin_cliente.get(f"/api/oficina/ciclos/{junta.ciclo.id}/cierre", secure=True)
    assert r.status_code == 200, r.content
    datos = r.json()
    assert datos["totales"]["ahorro"] == 30_000
    [fila] = datos["detalle"]
    assert fila["acciones"] == 3
    assert fila["deuda_pendiente"] > 0
    assert datos["politica_confirmada"] is False
