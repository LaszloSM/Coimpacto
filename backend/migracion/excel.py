"""Importación de los Excel mensuales de las ACAP que prepara CoimpactoB.

Cada libro tiene:
    "Listado de Miembros": encabezado con el reglamento del ciclo (duración, fecha de inicio,
        valor de la acción, fondo social, multa, % de interés) y una fila por miembro
        (#, nombres, cédula, celular, cargo).
    una hoja por reunión ("AGOSTO 2025"…): nombre de la ACAP, fecha, número de reunión y una
        fila por asistente (# acciones, $ ahorro, $ aporte social, $ multa).

Los encabezados no siempre están en la misma celda (en una hoja vienen fecha y frecuencia
cruzadas), así que todo se busca por la etiqueta y la fecha se toma de la celda que sí es fecha.
Se lee primero (`leer`), se revisa el informe y solo con `aplicar` se escribe en el libro, con
UUID deterministas: importar dos veces el mismo archivo no duplica nada.
"""

import datetime
import re
import unicodedata
import uuid
from dataclasses import dataclass, field
from decimal import Decimal

from django.db import transaction
from django.utils import timezone

from acap.models import Acap, Cargo, Ciclo, Membresia, Reunion
from cuentas.models import DatosPersonales, Titular, Usuario
from libro import servicio as libro
from libro.models import EstadoMovimiento, Movimiento

ESPACIO = uuid.UUID("5d1c1b8e-2f4c-4b1a-9a57-6f6c0e4f2a10")

CARGOS = {
    "presidente": Cargo.Tipo.PRESIDENTE,
    "presidenta": Cargo.Tipo.PRESIDENTE,
    "registrador": Cargo.Tipo.REGISTRADOR,
    "registradora": Cargo.Tipo.REGISTRADOR,
    "portador de caja": Cargo.Tipo.PORTADOR_CAJA,
    "portadora de caja": Cargo.Tipo.PORTADOR_CAJA,
    "contador 1": Cargo.Tipo.CONTADOR_1,
    "contadora 1": Cargo.Tipo.CONTADOR_1,
    "contador 2": Cargo.Tipo.CONTADOR_2,
    "contadora 2": Cargo.Tipo.CONTADOR_2,
}


def _norma(texto) -> str:
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", texto).strip().lower()


def _pesos(valor) -> int:
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return 0
    try:
        return int(round(float(valor)))
    except (TypeError, ValueError):
        return 0


def _fecha(valor) -> datetime.date | None:
    if isinstance(valor, datetime.datetime):
        return valor.date()
    if isinstance(valor, datetime.date):
        return valor
    return None


def _texto_doc(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return re.sub(r"[^0-9A-Za-z]", "", str(valor))[:20]


@dataclass
class Miembro:
    numero: int
    nombres: str
    apellidos: str
    documento: str
    telefono: str
    cargo: str | None


@dataclass
class Fila:
    numero: int
    nombre: str
    acciones: int
    ahorro: int
    fondo_social: int
    multa: int


@dataclass
class Hoja:
    titulo: str
    fecha: datetime.date | None
    numero: int | None
    filas: list[Fila] = field(default_factory=list)


@dataclass
class Libro:
    acap: str
    fecha_inicio: datetime.date | None
    duracion_meses: int
    valor_accion: int
    valor_fondo_social: int
    valor_multa: int
    tasa_interes: Decimal
    miembros: list[Miembro]
    hojas: list[Hoja]
    avisos: list[str]


def _buscar(filas, etiqueta: str, *, hasta: int = 8):
    """Valores a la derecha de la primera celda cuyo texto empieza por la etiqueta."""
    for fila in filas[:hasta]:
        for j, celda in enumerate(fila):
            if _norma(celda).startswith(etiqueta):
                return [v for v in fila[j + 1 :] if v not in (None, "")]
    return []


def _primera_fecha(filas, hasta=6):
    for fila in filas[:hasta]:
        for celda in fila:
            if _fecha(celda):
                return _fecha(celda)
    return None


def leer(ruta: str, *, duracion_meses: int | None = None) -> Libro:
    import openpyxl

    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    avisos: list[str] = []
    hojas_libro = wb.worksheets
    lista = next((h for h in hojas_libro if _norma(h.title).startswith("listado")), None)
    if lista is None:
        raise ValueError("No está la hoja «Listado de Miembros».")
    filas = [list(f) for f in lista.iter_rows(values_only=True)]

    # Las opciones de duración vienen en celdas (6, 9, 12) y la elegida lleva una «X».
    opciones = _buscar(filas, "duracion")
    marcada = [
        int(m.group(1))
        for v in opciones
        if "x" in str(v).lower() and (m := re.search(r"(\d+)", str(v)))
    ]
    duraciones = [int(float(m.group(1))) for v in opciones if (m := re.search(r"(\d+)", str(v)))]
    if duracion_meses is None:
        if marcada:
            duracion_meses = marcada[0]
        else:
            duracion_meses = max(duraciones) if duraciones else 12
            avisos.append(
                f"La duración del ciclo no está marcada con X entre {duraciones}; se usa "
                f"{duracion_meses} meses (cámbiela con --duracion)."
            )
    fecha_inicio = next((_fecha(v) for v in _buscar(filas, "fecha inicio") if _fecha(v)), None)
    valor_accion = _pesos((_buscar(filas, "valor de la accion") or [0])[0])
    fondo = _pesos((_buscar(filas, "valor fondo social") or [0])[0])
    multa = _pesos((_buscar(filas, "valor de la multa") or [0])[0])
    interes = (_buscar(filas, "% interes") or [0])[0]
    tasa = Decimal(str(interes or 0))
    if tasa < 1:  # 0.05 = 5 % mensual
        tasa = tasa * 100
    tasa = tasa.quantize(Decimal("0.001"))

    miembros: list[Miembro] = []
    inicio = next(
        (i for i, f in enumerate(filas) if f and _norma(f[0]) == "#"),
        None,
    )
    if inicio is None:
        raise ValueError("No se encontró la tabla de miembros (columna «#»).")
    for f in filas[inicio + 1 :]:
        if not f or len(f) < 2 or not str(f[1] or "").strip():
            continue
        try:
            numero = int(float(f[0]))
        except (TypeError, ValueError):
            continue
        partes = str(f[1]).strip().split()
        nombres, apellidos = partes[0], " ".join(partes[1:])
        cargo_texto = _norma(f[4] if len(f) > 4 else "")
        cargo = CARGOS.get(cargo_texto)
        if cargo_texto and cargo is None and not cargo_texto.startswith("miembro"):
            avisos.append(f"Miembro {numero}: cargo «{f[4]}» no reconocido; queda como miembro.")
        miembros.append(
            Miembro(
                numero=numero,
                nombres=nombres.title(),
                apellidos=apellidos.title(),
                documento=_texto_doc(f[2] if len(f) > 2 else None),
                telefono=_texto_doc(f[3] if len(f) > 3 else None),
                cargo=cargo,
            )
        )

    nombre_acap = ""
    hojas: list[Hoja] = []
    for hoja in hojas_libro:
        if hoja is lista:
            continue
        filas = [list(f) for f in hoja.iter_rows(values_only=True)]
        nombre = (_buscar(filas, "nombre de la acap") or [""])[0]
        nombre_acap = nombre_acap or str(nombre).strip()
        fecha = _primera_fecha(filas)
        if fecha is None:
            avisos.append(f"«{hoja.title}»: sin fecha de reunión; se omite.")
            continue
        etiqueta_fecha = _buscar(filas, "fecha de reunion")
        if etiqueta_fecha and not _fecha(etiqueta_fecha[0]):
            avisos.append(f"«{hoja.title}»: la fecha estaba en otra celda; se usó {fecha}.")
        numero = None
        for v in _buscar(filas, "n reunion") or _buscar(filas, "no reunion") or []:
            if isinstance(v, int | float):
                numero = int(v)
                break
        cabecera = next(
            (
                i
                for i, f in enumerate(filas)
                if len(f) > 1 and _norma(f[1]).startswith("nombre del")
            ),
            None,
        )
        if cabecera is None:
            avisos.append(f"«{hoja.title}»: sin tabla de asistentes; se omite.")
            continue
        registro = Hoja(titulo=hoja.title, fecha=fecha, numero=numero)
        for f in filas[cabecera + 1 :]:
            f = f + [None] * (6 - len(f))
            try:
                num = int(float(f[0]))
            except (TypeError, ValueError):
                continue
            if not str(f[1] or "").strip():
                continue
            acciones = _pesos(f[2])
            fila = Fila(
                numero=num,
                nombre=str(f[1]).strip(),
                acciones=acciones,
                ahorro=_pesos(f[3]),
                fondo_social=_pesos(f[4]),
                multa=_pesos(f[5]),
            )
            if valor_accion and fila.ahorro != acciones * valor_accion:
                avisos.append(
                    f"«{hoja.title}» miembro {num}: {acciones} acciones pero ahorro "
                    f"${fila.ahorro:,}; se registra por acciones (${acciones * valor_accion:,})."
                )
            if multa and fila.multa % multa:
                avisos.append(
                    f"«{hoja.title}» miembro {num}: multa ${fila.multa:,} no es múltiplo de "
                    f"${multa:,}; se omite."
                )
            registro.filas.append(fila)
        hojas.append(registro)

    hojas.sort(key=lambda h: h.fecha)
    numeros = {m.numero for m in miembros}
    for h in hojas:
        for f in h.filas:
            if f.numero not in numeros:
                avisos.append(f"«{h.titulo}»: el asistente {f.numero} no está en el listado.")
    return Libro(
        acap=nombre_acap or "ACAP sin nombre",
        fecha_inicio=fecha_inicio or (hojas[0].fecha if hojas else None),
        duracion_meses=duracion_meses,
        valor_accion=valor_accion,
        valor_fondo_social=fondo,
        valor_multa=multa,
        tasa_interes=tasa,
        miembros=miembros,
        hojas=hojas,
        avisos=avisos,
    )


def _sumar_meses(fecha: datetime.date, meses: int) -> datetime.date:
    total = fecha.month - 1 + meses
    anio, mes = fecha.year + total // 12, total % 12 + 1
    return datetime.date(anio, mes, min(fecha.day, 28))


def ciclos_de(datos: Libro, primer_numero: int) -> list[dict]:
    """Reparte las reuniones en ciclos por fecha: lo que pasa del fin del ciclo abre otro."""
    ciclos: list[dict] = []
    inicio = datos.fecha_inicio
    numero = primer_numero
    for h in datos.hojas:
        while inicio and h.fecha >= _sumar_meses(inicio, datos.duracion_meses):
            inicio = _sumar_meses(inicio, datos.duracion_meses)
            numero += 1
        if not ciclos or ciclos[-1]["numero"] != numero:
            ciclos.append({"numero": numero, "inicio": inicio, "hojas": []})
        ciclos[-1]["hojas"].append(h)
    return ciclos


def _id(*partes) -> uuid.UUID:
    return uuid.uuid5(ESPACIO, "|".join(str(p) for p in partes))


@transaction.atomic
def aplicar(
    datos: Libro,
    *,
    usuario: Usuario,
    primer_ciclo: int = 1,
    zona: str = "",
    departamento: str = "",
    municipio: str = "",
) -> dict:
    """Escribe ACAP, ciclos, miembros, cargos, reuniones cerradas y movimientos confirmados."""
    acap, _ = Acap.objects.get_or_create(nombre=datos.acap)
    for campo, valor in (("zona", zona), ("departamento", departamento), ("municipio", municipio)):
        if valor:
            setattr(acap, campo, valor)
    acap.save()

    titulares: dict[int, Titular] = {}
    for m in datos.miembros:
        membresia = Membresia.objects.filter(acap=acap, numero_miembro=m.numero).first()
        if membresia is None:
            titular = Titular.objects.create()
            membresia = Membresia.objects.create(
                acap=acap,
                titular=titular,
                numero_miembro=m.numero,
                fecha_ingreso=datos.fecha_inicio or timezone.localdate(),
            )
        # Datos personales migrados sin consentimiento: no se muestran hasta capturarlo (T8).
        DatosPersonales.objects.update_or_create(
            titular_id=membresia.titular_id,
            defaults={
                "nombres": m.nombres[:80],
                "apellidos": m.apellidos[:80],
                "documento": m.documento,
                "telefono": m.telefono,
            },
        )
        titulares[m.numero] = membresia.titular

    max_acciones = max([f.acciones for h in datos.hojas for f in h.filas] + [5])
    cuenta = {"ciclos": 0, "reuniones": 0, "movimientos": 0}
    plan = ciclos_de(datos, primer_ciclo)
    for i, c in enumerate(plan):
        ultimo = i == len(plan) - 1
        ciclo, creado = Ciclo.objects.get_or_create(
            acap=acap,
            numero=c["numero"],
            defaults={
                "fecha_inicio": c["inicio"] or c["hojas"][0].fecha,
                "duracion_meses": datos.duracion_meses,
                "valor_accion": datos.valor_accion,
                "max_acciones_reunion": max_acciones,
                "valor_fondo_social": datos.valor_fondo_social,
                "valor_multa": datos.valor_multa,
                "tasa_interes": datos.tasa_interes,
                "estado": Ciclo.Estado.CERRADO,
            },
        )
        cuenta["ciclos"] += int(creado)
        if ultimo and ciclo.estado != Ciclo.Estado.ACTIVO:
            Ciclo.objects.filter(acap=acap, estado=Ciclo.Estado.ACTIVO).exclude(pk=ciclo.pk).update(
                estado=Ciclo.Estado.CERRADO
            )
            ciclo.estado = Ciclo.Estado.ACTIVO
            ciclo.save(update_fields=["estado"])
        if ultimo:
            for m in datos.miembros:
                if (
                    m.cargo
                    and not Cargo.objects.filter(ciclo=ciclo, tipo=m.cargo, vigente=True).exists()
                ):
                    membresia = Membresia.objects.get(acap=acap, numero_miembro=m.numero)
                    if not Cargo.objects.filter(
                        ciclo=ciclo, membresia=membresia, vigente=True
                    ).exists():
                        Cargo.objects.create(ciclo=ciclo, membresia=membresia, tipo=m.cargo)

        for n, h in enumerate(c["hojas"], start=1):
            rid = _id(acap.id, "reunion", h.fecha.isoformat())
            reunion = Reunion.objects.filter(id=rid).first()
            if reunion is None:
                numero = n
                while Reunion.objects.filter(ciclo=ciclo, numero=numero).exists():
                    numero += 1
                reunion = Reunion.abrir(id=rid, ciclo=ciclo, numero=numero, fecha=h.fecha)
                cuenta["reuniones"] += 1
            elif reunion.estado == Reunion.Estado.CERRADA:
                continue
            for f in h.filas:
                titular = titulares.get(f.numero)
                if titular is None:
                    continue
                movs = []
                if f.acciones > 0:
                    movs.append(
                        libro.registrar_ahorro(
                            id=_id(rid, f.numero, "ahorro"),
                            reunion=reunion,
                            titular=titular,
                            acciones=f.acciones,
                            actor=usuario,
                            dispositivo="importe-excel",
                        )
                    )
                if f.fondo_social > 0:
                    movs.append(
                        libro.registrar_movimiento(
                            id=_id(rid, f.numero, "fondo"),
                            reunion=reunion,
                            tipo=Movimiento.Tipo.FONDO_SOCIAL,
                            monto=f.fondo_social,
                            titular=titular,
                            actor=usuario,
                            dispositivo="importe-excel",
                        )
                    )
                if f.multa > 0 and datos.valor_multa and f.multa % datos.valor_multa == 0:
                    movs.append(
                        libro.registrar_multas(
                            id=_id(rid, f.numero, "multa"),
                            reunion=reunion,
                            titular=titular,
                            cantidad=f.multa // datos.valor_multa,
                            motivo="Importado del Excel",
                            actor=usuario,
                            dispositivo="importe-excel",
                        )
                    )
                for mov in movs:
                    if libro.estado_actual(mov) == EstadoMovimiento.PROVISIONAL:
                        libro.cambiar_estado(
                            mov,
                            EstadoMovimiento.CONFIRMADO,
                            actor=usuario,
                            codigo="IMPORTE_EXCEL",
                        )
                        cuenta["movimientos"] += 1
            reunion.estado = Reunion.Estado.CERRADA
            reunion.cerrada_en = timezone.make_aware(
                datetime.datetime.combine(h.fecha, datetime.time(18, 0))
            )
            reunion.save(update_fields=["estado", "cerrada_en"])
    return {"acap": str(acap.id), **cuenta}
