"""Módulos de análisis del panel de la Fundación: Ahorro, Crédito, Cartera, Análisis y Reportes.

Todo sale del libro: movimientos vigentes (sin rechazados ni en conflicto) con sus correcciones
(cada reverso cuenta en el tipo del movimiento que corrige). Sin datos personales: los
préstamos se identifican por ACAP y número de miembro (Ley 1581). Las reglas de mora son las
del tablero (acap/tablero.py).
"""

import datetime
import io
from collections import defaultdict
from statistics import mean

from django.utils import timezone

from acap.models import Acap, Ciclo, Membresia, Reunion
from acap.tablero import _estado_prestamo
from libro.models import EstadoMovimiento, Movimiento, Prestamo

E = EstadoMovimiento
T = Movimiento.Tipo
TIPOS_AHORRO = (T.AHORRO, T.FONDO_SOCIAL, T.MULTA)
TRAMOS = [
    ("al_dia", "Al día", 0, 0),
    ("1_30", "1 a 30 días", 1, 30),
    ("31_60", "31 a 60 días", 31, 60),
    ("61_90", "61 a 90 días", 61, 90),
    ("mas_90", "Más de 90 días", 91, 10**6),
]


def _efectivos(filtro=None):
    """(tipo efectivo, monto, reunión, titular, préstamo) de cada movimiento vigente."""
    qs = Movimiento.objects.con_estado().exclude(estado__in=[E.RECHAZADO, E.EN_CONFLICTO])
    if filtro is not None:
        qs = qs.filter(filtro)
    qs = qs.select_related("reunion", "reunion__ciclo", "corrige")
    for m in qs:
        tipo = m.corrige.tipo if m.tipo == T.REVERSO and m.corrige else m.tipo
        yield {
            "tipo": tipo,
            "monto": m.monto,
            "reunion": m.reunion,
            "ciclo_id": m.reunion.ciclo_id,
            "acap_id": m.reunion.ciclo.acap_id,
            "titular_id": m.titular_id,
            "prestamo_id": m.prestamo_id,
            "detalle": m.detalle,
        }


def _mes(fecha: datetime.date) -> str:
    return fecha.isoformat()[:7]


MESES = [
    "enero",
    "febrero",
    "marzo",
    "abril",
    "mayo",
    "junio",
    "julio",
    "agosto",
    "septiembre",
    "octubre",
    "noviembre",
    "diciembre",
]


def _mes_texto(aaaa_mm: str) -> str:
    """ "2026-10" → "octubre de 2026"."""
    return f"{MESES[int(aaaa_mm[5:7]) - 1]} de {aaaa_mm[:4]}"


def _ultimos_meses(hoy: datetime.date, n: int) -> list[str]:
    meses, fecha = [], hoy.replace(day=1)
    for _ in range(n):
        meses.append(_mes(fecha))
        fecha = (fecha - datetime.timedelta(days=1)).replace(day=1)
    return list(reversed(meses))


def _acaps():
    return {a.id: a for a in Acap.objects.order_by("nombre")}


def _ciclos_activos():
    return {c.acap_id: c for c in Ciclo.objects.filter(estado=Ciclo.Estado.ACTIVO)}


# ---------------------------------------------------------------- Ahorro


def ahorro(hoy: datetime.date | None = None) -> dict:
    hoy = hoy or timezone.localdate()
    acaps, activos = _acaps(), _ciclos_activos()
    por_acap = defaultdict(lambda: defaultdict(int))
    historico = defaultdict(int)
    por_mes = defaultdict(lambda: defaultdict(int))
    ahorradores = defaultdict(set)  # reunión -> titulares que ahorraron
    for m in _efectivos():
        if m["tipo"] not in TIPOS_AHORRO:
            continue
        historico[m["tipo"]] += m["monto"]
        por_mes[_mes(m["reunion"].fecha)][m["tipo"]] += m["monto"]
        if activos.get(m["acap_id"]) and m["ciclo_id"] == activos[m["acap_id"]].id:
            por_acap[m["acap_id"]][m["tipo"]] += m["monto"]
        if m["tipo"] == T.AHORRO and m["monto"] > 0:
            ahorradores[m["reunion"].id].add(m["titular_id"])

    miembros = defaultdict(int)
    for acap_id in Membresia.objects.filter(activa=True).values_list("acap_id", flat=True):
        miembros[acap_id] += 1
    ultimas = {}
    for r in Reunion.objects.select_related("ciclo").order_by("ciclo__acap_id", "-fecha"):
        ultimas.setdefault(r.ciclo.acap_id, r)

    filas = []
    for acap_id, acap in acaps.items():
        ciclo = activos.get(acap_id)
        datos = por_acap.get(acap_id, {})
        ultima = ultimas.get(acap_id)
        n = miembros.get(acap_id, 0)
        participacion = (
            round(100 * len(ahorradores.get(ultima.id, set())) / n, 1) if ultima and n else None
        )
        reuniones = ciclo.reuniones.count() if ciclo else 0
        filas.append(
            {
                "id": str(acap_id),
                "nombre": acap.nombre,
                "zona": acap.zona or "Sin zona",
                "ciclo": ciclo.numero if ciclo else None,
                "reuniones": reuniones,
                "miembros": n,
                "ahorro": datos.get(T.AHORRO, 0),
                "fondo_social": datos.get(T.FONDO_SOCIAL, 0),
                "multas": datos.get(T.MULTA, 0),
                "por_miembro": round(datos.get(T.AHORRO, 0) / n) if n else 0,
                "meta": ciclo.meta_ahorro if ciclo else None,
                "ultima_reunion": ultima.fecha.isoformat() if ultima else None,
                "participacion": participacion,
            }
        )
    filas.sort(key=lambda f: -f["ahorro"])

    zonas = defaultdict(lambda: {"ahorro": 0, "comunidades": 0, "miembros": 0})
    for f in filas:
        z = zonas[f["zona"]]
        z["ahorro"] += f["ahorro"]
        z["comunidades"] += 1
        z["miembros"] += f["miembros"]

    total_ciclo = sum(f["ahorro"] for f in filas)
    total_miembros = sum(f["miembros"] for f in filas)
    participaciones = [f["participacion"] for f in filas if f["participacion"] is not None]
    return {
        "resumen": {
            "ahorro_ciclo": total_ciclo,
            "fondo_social_ciclo": sum(f["fondo_social"] for f in filas),
            "multas_ciclo": sum(f["multas"] for f in filas),
            "ahorro_historico": historico[T.AHORRO],
            "por_miembro": round(total_ciclo / total_miembros) if total_miembros else 0,
            "participacion": round(mean(participaciones), 1) if participaciones else None,
        },
        "comunidades": filas,
        "zonas": [
            {"zona": z, **v} for z, v in sorted(zonas.items(), key=lambda kv: -kv[1]["ahorro"])
        ],
        "mensual": [
            {
                "mes": mes,
                "ahorro": por_mes[mes].get(T.AHORRO, 0),
                "fondo_social": por_mes[mes].get(T.FONDO_SOCIAL, 0),
                "multas": por_mes[mes].get(T.MULTA, 0),
            }
            for mes in _ultimos_meses(hoy, 12)
        ],
    }


# ---------------------------------------------------------------- Crédito y cartera


def _prestamos(corte: datetime.date) -> list[dict]:
    acaps = _acaps()
    numeros = {
        (m.acap_id, m.titular_id): m.numero_miembro
        for m in Membresia.objects.only("acap_id", "titular_id", "numero_miembro")
    }
    filas = []
    for p in Prestamo.objects.select_related("reunion", "ciclo").order_by("-reunion__fecha"):
        e = _estado_prestamo(p, corte)
        abonado = p.total_pagar - e["saldo"]
        dias = 0
        if e["atraso"] > 0 and p.cuota_mensual:
            cuotas_pagadas = abonado // p.cuota_mensual
            vence = p.reunion.fecha + datetime.timedelta(days=30 * (cuotas_pagadas + 1))
            dias = max(0, (corte - vence).days)
        acap = acaps.get(p.ciclo.acap_id)
        filas.append(
            {
                "id": str(p.id),
                "acap_id": str(p.ciclo.acap_id),
                "acap": acap.nombre if acap else "",
                "zona": (acap.zona if acap else "") or "Sin zona",
                "miembro": numeros.get((p.ciclo.acap_id, p.titular_id)),
                "fecha": p.reunion.fecha.isoformat(),
                "capital": p.capital,
                "tasa": float(p.tasa_interes),
                "cuotas": p.num_cuotas,
                "cuota": p.cuota_mensual,
                "total": p.total_pagar,
                "abonado": abonado,
                "saldo": e["saldo"],
                "debidas": e["debidas"],
                "atraso": e["atraso"],
                "dias_atraso": dias,
                "estado": "pagado"
                if e["saldo"] <= 0
                else ("mora" if e["atraso"] > 0 else "al_dia"),
            }
        )
    return filas


def credito(hoy: datetime.date | None = None) -> dict:
    hoy = hoy or timezone.localdate()
    prestamos = _prestamos(hoy)
    por_mes = defaultdict(lambda: {"desembolsado": 0, "abonado": 0, "prestamos": 0})
    for m in _efectivos():
        mes = _mes(m["reunion"].fecha)
        if m["tipo"] == T.DESEMBOLSO:
            por_mes[mes]["desembolsado"] += m["monto"]
        elif m["tipo"] == T.ABONO:
            por_mes[mes]["abonado"] += m["monto"]
    for p in prestamos:
        por_mes[p["fecha"][:7]]["prestamos"] += 1

    por_acap = defaultdict(lambda: {"prestamos": 0, "capital": 0, "intereses": 0, "vigentes": 0})
    for p in prestamos:
        a = por_acap[(p["acap_id"], p["acap"])]
        a["prestamos"] += 1
        a["capital"] += p["capital"]
        a["intereses"] += p["total"] - p["capital"]
        a["vigentes"] += int(p["saldo"] > 0)

    rangos = [
        ("Hasta $200 mil", 0, 200_000),
        ("$200 mil a $500 mil", 200_001, 500_000),
        ("$500 mil a $1 M", 500_001, 1_000_000),
        ("Más de $1 M", 1_000_001, 10**15),
    ]
    capital = [p["capital"] for p in prestamos]
    return {
        "resumen": {
            "prestamos": len(prestamos),
            "vigentes": sum(1 for p in prestamos if p["saldo"] > 0),
            "desembolsado": sum(capital),
            "intereses_pactados": sum(p["total"] - p["capital"] for p in prestamos),
            "capital_promedio": round(mean(capital)) if capital else 0,
            "plazo_promedio": round(mean(p["cuotas"] for p in prestamos), 1) if prestamos else 0,
            "tasa_promedio": round(mean(p["tasa"] for p in prestamos), 2) if prestamos else 0,
        },
        "comunidades": sorted(
            [{"id": k[0], "nombre": k[1], **v} for k, v in por_acap.items()],
            key=lambda f: -f["capital"],
        ),
        "rangos": [
            {"rango": r, "prestamos": sum(1 for c in capital if lo <= c <= hi)}
            for r, lo, hi in rangos
        ],
        "mensual": [{"mes": mes, **por_mes[mes]} for mes in _ultimos_meses(hoy, 12)],
        "recientes": prestamos[:20],
    }


def cartera(hoy: datetime.date | None = None) -> dict:
    hoy = hoy or timezone.localdate()
    vigentes = [p for p in _prestamos(hoy) if p["saldo"] > 0]
    tramos = []
    for clave, texto, lo, hi in TRAMOS:
        en = [p for p in vigentes if lo <= p["dias_atraso"] <= hi]
        if clave == "al_dia":
            en = [p for p in vigentes if p["atraso"] == 0]
        elif lo <= 30:
            en = [p for p in en if p["atraso"] > 0]
        tramos.append(
            {
                "tramo": clave,
                "texto": texto,
                "prestamos": len(en),
                "saldo": sum(p["saldo"] for p in en),
            }
        )
    activa = sum(p["saldo"] for p in vigentes)
    vencida = sum(p["saldo"] for p in vigentes if p["atraso"] > 0)
    por_acap = defaultdict(lambda: {"saldo": 0, "vencida": 0, "prestamos": 0, "en_mora": 0})
    for p in vigentes:
        a = por_acap[(p["acap_id"], p["acap"], p["zona"])]
        a["saldo"] += p["saldo"]
        a["prestamos"] += 1
        if p["atraso"] > 0:
            a["vencida"] += p["saldo"]
            a["en_mora"] += 1
    comunidades = [
        {
            "id": k[0],
            "nombre": k[1],
            "zona": k[2],
            **v,
            "mora": round(100 * v["vencida"] / v["saldo"], 1) if v["saldo"] else 0.0,
        }
        for k, v in por_acap.items()
    ]
    return {
        "resumen": {
            "activa": activa,
            "vencida": vencida,
            "mora": round(100 * vencida / activa, 1) if activa else 0.0,
            "prestamos": len(vigentes),
            "en_mora": sum(1 for p in vigentes if p["atraso"] > 0),
            "atraso_total": sum(p["atraso"] for p in vigentes),
        },
        "tramos": tramos,
        "comunidades": sorted(comunidades, key=lambda f: -f["mora"]),
        "prestamos": sorted(vigentes, key=lambda p: (-p["dias_atraso"], -p["saldo"])),
    }


# ---------------------------------------------------------------- Análisis


def analisis(hoy: datetime.date | None = None) -> dict:
    """Lecturas por reglas sobre el libro (sin IA en el MVP): qué cambia y dónde mirar."""
    hoy = hoy or timezone.localdate()
    a = ahorro(hoy)
    c = cartera(hoy)
    lecturas = []

    meses = [m for m in a["mensual"] if m["ahorro"]]
    if len(meses) >= 2:
        ultimo, previo = meses[-1], meses[-2]
        cambio = round(100 * (ultimo["ahorro"] - previo["ahorro"]) / previo["ahorro"], 1)
        lecturas.append(
            {
                "tono": "bien" if cambio >= 0 else "alerta",
                "titulo": "Ritmo del ahorro",
                "texto": f"En {_mes_texto(ultimo['mes'])} se ahorraron "
                f"${ultimo['ahorro']:,}".replace(",", ".")
                + f" ({'+' if cambio >= 0 else ''}{str(cambio).replace('.', ',')} % frente a "
                f"{_mes_texto(previo['mes'])}).",
            }
        )

    for f in a["comunidades"]:
        if f["participacion"] is not None and f["participacion"] < 60:
            lecturas.append(
                {
                    "tono": "alerta",
                    "titulo": f"Participación en {f['nombre']}",
                    "acap_id": f["id"],
                    "texto": f"Solo {f['participacion']:.0f} % de los miembros ahorró en la "
                    "última reunión. Conviene preguntar a la junta qué pasa.",
                }
            )
        if f["ultima_reunion"]:
            dias = (hoy - datetime.date.fromisoformat(f["ultima_reunion"])).days
            if dias > 45:
                lecturas.append(
                    {
                        "tono": "neutro",
                        "titulo": f"{f['nombre']} sin reunión registrada",
                        "acap_id": f["id"],
                        "texto": f"La última reunión registrada fue hace {dias} días.",
                    }
                )
        if f["meta"]:
            avance = round(100 * f["ahorro"] / f["meta"])
            lecturas.append(
                {
                    "tono": "bien" if avance >= 80 else "neutro",
                    "titulo": f"Meta de {f['nombre']}",
                    "acap_id": f["id"],
                    "texto": f"Lleva {avance} % de su meta de ahorro del ciclo.",
                }
            )

    if c["resumen"]["activa"]:
        top = max(c["comunidades"], key=lambda f: f["saldo"])
        parte = round(100 * top["saldo"] / c["resumen"]["activa"])
        if parte >= 40 and len(c["comunidades"]) > 1:
            lecturas.append(
                {
                    "tono": "neutro",
                    "titulo": "Concentración de la cartera",
                    "acap_id": top["id"],
                    "texto": f"{top['nombre']} tiene el {parte} % de toda la cartera activa.",
                }
            )
        for f in c["comunidades"]:
            if f["mora"] > 10:
                lecturas.append(
                    {
                        "tono": "alerta",
                        "titulo": f"Mora en {f['nombre']}",
                        "acap_id": f["id"],
                        "texto": f"{f['mora']:.1f} % de su cartera está en mora "
                        f"({f['en_mora']} de {f['prestamos']} préstamos).",
                    }
                )

    # Ahorro por miembro de cada comunidad frente al promedio: quién va adelante.
    con_miembros = [f for f in a["comunidades"] if f["miembros"]]
    if len(con_miembros) >= 2:
        promedio = mean(f["por_miembro"] for f in con_miembros)
        lider = max(con_miembros, key=lambda f: f["por_miembro"])
        if promedio and lider["por_miembro"] > promedio * 1.2:
            lecturas.append(
                {
                    "tono": "bien",
                    "titulo": "Comunidad que va adelante",
                    "acap_id": lider["id"],
                    "texto": f"En {lider['nombre']} cada miembro ahorró en promedio "
                    f"${lider['por_miembro']:,}, por encima del promedio.".replace(",", "."),
                }
            )

    orden = {"alerta": 0, "neutro": 1, "bien": 2}
    lecturas.sort(key=lambda lectura: orden[lectura["tono"]])
    return {
        "generado_en": timezone.now().isoformat(),
        "lecturas": lecturas,
        "indicadores": {
            "participacion": a["resumen"]["participacion"],
            "por_miembro": a["resumen"]["por_miembro"],
            "mora": c["resumen"]["mora"],
            "comunidades_sin_reunion": sum(
                1
                for f in a["comunidades"]
                if not f["ultima_reunion"]
                or (hoy - datetime.date.fromisoformat(f["ultima_reunion"])).days > 45
            ),
        },
    }


# ---------------------------------------------------------------- Reportes


def _hoja(wb, titulo: str, cabecera: list[str], filas: list[list]):
    from openpyxl.styles import Font, PatternFill

    ws = wb.create_sheet(titulo)
    ws.append(cabecera)
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1C1C1C")
    for fila in filas:
        ws.append(fila)
    for i, nombre in enumerate(cabecera, start=1):
        largo = max([len(str(nombre))] + [len(str(f[i - 1] or "")) for f in filas] + [8])
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(largo + 2, 40)
    ws.freeze_panes = "A2"
    return ws


def reporte_consolidado(hoy: datetime.date | None = None) -> bytes:
    """Libro de Excel con el consolidado de la Fundación (sin datos personales)."""
    from openpyxl import Workbook

    hoy = hoy or timezone.localdate()
    a, cr, ca = ahorro(hoy), credito(hoy), cartera(hoy)
    wb = Workbook()
    wb.remove(wb.active)
    r = a["resumen"]
    _hoja(
        wb,
        "Resumen",
        ["Indicador", "Valor"],
        [
            ["Fecha de corte", hoy.isoformat()],
            ["Ahorro del ciclo (acciones)", r["ahorro_ciclo"]],
            ["Fondo social del ciclo", r["fondo_social_ciclo"]],
            ["Multas del ciclo", r["multas_ciclo"]],
            ["Ahorro histórico", r["ahorro_historico"]],
            ["Ahorro promedio por miembro", r["por_miembro"]],
            ["Participación en la última reunión (%)", r["participacion"]],
            ["Préstamos otorgados", cr["resumen"]["prestamos"]],
            ["Capital desembolsado", cr["resumen"]["desembolsado"]],
            ["Cartera activa", ca["resumen"]["activa"]],
            ["Cartera vencida", ca["resumen"]["vencida"]],
            ["Índice de mora (%)", ca["resumen"]["mora"]],
        ],
    )
    _hoja(
        wb,
        "Ahorro por ACAP",
        [
            "ACAP",
            "Zona",
            "Ciclo",
            "Reuniones",
            "Miembros",
            "Ahorro",
            "Fondo social",
            "Multas",
            "Ahorro por miembro",
            "Meta",
            "Última reunión",
            "Participación %",
        ],
        [
            [
                f["nombre"],
                f["zona"],
                f["ciclo"],
                f["reuniones"],
                f["miembros"],
                f["ahorro"],
                f["fondo_social"],
                f["multas"],
                f["por_miembro"],
                f["meta"],
                f["ultima_reunion"],
                f["participacion"],
            ]
            for f in a["comunidades"]
        ],
    )
    _hoja(
        wb,
        "Ahorro mensual",
        ["Mes", "Ahorro", "Fondo social", "Multas"],
        [[m["mes"], m["ahorro"], m["fondo_social"], m["multas"]] for m in a["mensual"]],
    )
    _hoja(
        wb,
        "Cartera por ACAP",
        ["ACAP", "Zona", "Préstamos vigentes", "En mora", "Saldo", "Vencida", "Mora %"],
        [
            [
                f["nombre"],
                f["zona"],
                f["prestamos"],
                f["en_mora"],
                f["saldo"],
                f["vencida"],
                f["mora"],
            ]
            for f in ca["comunidades"]
        ],
    )
    _hoja(
        wb,
        "Préstamos vigentes",
        [
            "ACAP",
            "Miembro #",
            "Fecha",
            "Capital",
            "Tasa %",
            "Cuotas",
            "Cuota",
            "Abonado",
            "Saldo",
            "Atraso",
            "Días de atraso",
        ],
        [
            [
                p["acap"],
                p["miembro"],
                p["fecha"],
                p["capital"],
                p["tasa"],
                p["cuotas"],
                p["cuota"],
                p["abonado"],
                p["saldo"],
                p["atraso"],
                p["dias_atraso"],
            ]
            for p in ca["prestamos"]
        ],
    )
    reuniones = Reunion.objects.select_related("ciclo__acap").order_by("-fecha")
    totales = defaultdict(lambda: defaultdict(int))
    for m in _efectivos():
        totales[m["reunion"].id][m["tipo"]] += m["monto"]
    _hoja(
        wb,
        "Reuniones",
        [
            "ACAP",
            "Ciclo",
            "Reunión",
            "Fecha",
            "Estado",
            "Ahorro",
            "Fondo social",
            "Multas",
            "Desembolsos",
            "Abonos",
        ],
        [
            [
                r.ciclo.acap.nombre,
                r.ciclo.numero,
                r.numero,
                r.fecha.isoformat(),
                r.get_estado_display(),
                totales[r.id].get(T.AHORRO, 0),
                totales[r.id].get(T.FONDO_SOCIAL, 0),
                totales[r.id].get(T.MULTA, 0),
                totales[r.id].get(T.DESEMBOLSO, 0),
                totales[r.id].get(T.ABONO, 0),
            ]
            for r in reuniones
        ],
    )
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
